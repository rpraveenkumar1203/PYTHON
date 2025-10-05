from __future__ import annotations
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Dict
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from ..config import XLSX_PATH

COLUMNS = ["EntryID", "Name", "Amount", "Currency", "CreatedAt", "ReceiptFilename"]


def ensure_workbook() -> None:
    if not XLSX_PATH.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Entries"
        ws.append(COLUMNS)
        widths = [18, 24, 12, 10, 24, 40]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        wb.save(XLSX_PATH)


def read_all_rows() -> List[Dict]:
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb.active
    data: List[Dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        data.append(dict(zip(COLUMNS, row)))
    wb.close()
    return data


def next_entry_id(now: datetime | None = None) -> str:
    ensure_workbook()
    try:
        now = now or datetime.now()
        prefix = f"R-{now.strftime('%Y%m%d')}"
        wb = load_workbook(XLSX_PATH, read_only=True)
        ws = wb.active
        max_seq = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            entry_id = r[0]
            if isinstance(entry_id, str) and entry_id.startswith(prefix + "-"):
                m = re.search(r"(\d+)$", entry_id)
                if m:
                    max_seq = max(max_seq, int(m.group(1)))
        wb.close()
        return f"{prefix}-{max_seq + 1:04d}"
    except Exception:
        return f"R-{uuid.uuid4().hex[:8].upper()}"


def append_row(entry: dict) -> None:
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb.active
    ws.append([entry[c] for c in COLUMNS])
    wb.save(XLSX_PATH)
    wb.close()
